def get_fields_by_role(field_dict, role):
    return {key: value for (key, value) in field_dict.items() if value == role}


def write_to_file(filepath, data):
    with open(filepath, "w") as f:
        for k, v in data.items():
            f.write(str(k) + " " + str(v) + "\n")


def generate_pdf(pdf_input, pdf_output, data):
    import os
    os.system("C:\\Data\\Projects\\Practica\\practicaServer\\controller\\editor\\FormEditor.exe " + pdf_input + " " + pdf_output + " " + data)


def fill_docx_activity(docx_input, docx_output, activity):
    from docx import Document
    doc = Document(docx_input)
    table = doc.tables[0]
    for i in range(len(activity)):
        table.add_row()
        table.cell(i + 1, 0).text = activity[i][0] + "; " + activity[i][1]
        table.cell(i + 1, 1).text = activity[i][2]
    doc.save(docx_output)


fill_docx_activity("RaportActivitatePractica.docx", "o.docx", [("a", "b", "c"), ("d", "e", "f"), ("g", "h", "i")])

